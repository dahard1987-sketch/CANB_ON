from __future__ import annotations

import io
import hashlib
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

import app as web_app
from update_excel import inspect_excel_file, normalize_level_value


TEST_PASSWORD = "readi-test-password"
TEST_PASSWORD_HASH = hashlib.sha256(TEST_PASSWORD.encode("utf-8")).hexdigest()


def workbook_bytes(level_values: list[str], sheet_name: str = "학생 상세") -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(["번호", "이름", "", "", "", "", "레벨", "", "", "", ""])
    for index, level in enumerate(level_values, 1):
        worksheet.append([index, f"학생 {index}", "", "", "", "", level, "", "", "", ""])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class LevelDetectionTests(unittest.TestCase):
    def test_normalizes_supported_level_variants(self):
        self.assertEqual(normalize_level_value("nona3"), "Nona 3")
        self.assertEqual(normalize_level_value("레벨: HEPTA  1반"), "Hepta 1")
        self.assertEqual(normalize_level_value("hexa2"), "Hexa 2")
        self.assertIsNone(normalize_level_value("Hepta 3"))
        self.assertIsNone(normalize_level_value("Hexa 3"))
        self.assertIsNone(normalize_level_value("레벨"))

    def test_reads_level_from_column_g(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "filename-does-not-matter.xlsx"
            path.write_bytes(workbook_bytes(["Octa 2", "Octa 2"]))
            level, rows = inspect_excel_file(path)
        self.assertEqual(level, "Octa 2")
        self.assertEqual(len(rows), 3)

    def test_rejects_multiple_levels(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mixed.xlsx"
            path.write_bytes(workbook_bytes(["Hepta 1", "Nona 3"]))
            with self.assertRaisesRegex(ValueError, "레벨을 하나로 확인"):
                inspect_excel_file(path)

    def test_ignores_one_stray_level_when_one_level_is_dominant(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mostly-hexa.xlsx"
            path.write_bytes(workbook_bytes(["Hexa 1"] * 9 + ["Octa 1"]))
            level, _rows = inspect_excel_file(path)
        self.assertEqual(level, "Hexa 1")


class AuthenticationTests(unittest.TestCase):
    def setUp(self):
        web_app.app.config.update(
            TESTING=True,
            APP_PASSWORD_SHA256=TEST_PASSWORD_HASH,
            SECRET_KEY="test-session-secret",
            SESSION_COOKIE_SECURE=False,
        )
        web_app._login_attempts.clear()
        self.client = web_app.app.test_client()

    def test_password_gate_login_and_logout(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.headers["Location"].endswith("/login"))

        wrong = self.client.post("/login", data={"password": "wrong"})
        self.assertEqual(wrong.status_code, 401)
        self.assertIn("비밀번호가 맞지 않습니다", wrong.get_data(as_text=True))

        login = self.client.post("/login", data={"password": TEST_PASSWORD})
        self.assertEqual(login.status_code, 302)

        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertIn("READi 업데이터", page.get_data(as_text=True))

        logout = self.client.post("/logout")
        self.assertEqual(logout.status_code, 302)
        self.assertEqual(self.client.get("/").status_code, 302)

    def test_api_requires_login(self):
        response = self.client.post("/api/files")
        self.assertEqual(response.status_code, 401)
        self.assertIn("로그인이 만료", response.get_json()["error"])


class UploadApiTests(unittest.TestCase):
    def setUp(self):
        web_app.app.config.update(
            TESTING=True,
            APP_PASSWORD_SHA256=TEST_PASSWORD_HASH,
            SECRET_KEY="test-session-secret",
            SESSION_COOKIE_SECURE=False,
        )
        web_app._login_attempts.clear()
        self.client = web_app.app.test_client()
        self.client.post("/login", data={"password": TEST_PASSWORD})

    def test_multiple_files_are_classified_without_filename(self):
        response = self.client.post(
            "/api/files",
            data={
                "files": [
                    (io.BytesIO(workbook_bytes(["Hepta 2"])), "alpha.xlsx"),
                    (io.BytesIO(workbook_bytes(["Nona 1"])), "beta.xlsx"),
                ]
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        accepted = {item["level"]: item for item in payload["results"] if item["accepted"]}
        self.assertEqual(set(accepted), {"Hepta 2", "Nona 1"})
        self.assertEqual(accepted["Hepta 2"]["filename"], "alpha.xlsx")

    def test_invalid_workbook_is_reported_per_file(self):
        response = self.client.post(
            "/api/files",
            data={"files": (io.BytesIO(b"not an xlsx"), "broken.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        result = response.get_json()["results"][0]
        self.assertFalse(result["accepted"])
        self.assertTrue(result["error"])

    def test_publish_updates_detected_sheet(self):
        with (
            patch.object(web_app, "create_sheets_service", return_value=object()),
            patch.object(web_app, "get_sheet_names", return_value={"Octa 3", "Octa 3_표시"}),
            patch.object(web_app, "clear_target_sheet") as clear,
            patch.object(web_app, "upload_values", return_value=22) as upload,
            patch.object(web_app, "update_display_timestamp") as timestamp,
        ):
            response = self.client.post(
                "/api/publish",
                data={
                    "files": (
                        io.BytesIO(workbook_bytes(["Octa 3"])),
                        "anything.xlsx",
                    )
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["summary"], {"success": 1, "failed": 0})
        clear.assert_called_once_with(unittest.mock.ANY, "Octa 3")
        self.assertEqual(upload.call_args.args[1], "Octa 3")
        timestamp.assert_called_once()

    def test_missing_sheet_stops_before_any_clear(self):
        with (
            patch.object(web_app, "create_sheets_service", return_value=object()),
            patch.object(web_app, "get_sheet_names", return_value=set()),
            patch.object(web_app, "clear_target_sheet") as clear,
        ):
            response = self.client.post(
                "/api/publish",
                data={
                    "files": (
                        io.BytesIO(workbook_bytes(["Hexa 1"])),
                        "hexa.xlsx",
                    )
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 400)
        clear.assert_not_called()

    def test_cross_origin_write_is_rejected(self):
        response = self.client.post(
            "/api/files",
            headers={"Origin": "https://malicious.example"},
            data={"files": (io.BytesIO(workbook_bytes(["Nona 2"])), "nona.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 403)

    def test_health_check_and_security_headers(self):
        response = self.client.get("/healthz")
        self.assertEqual(response.get_json(), {"status": "ok"})
        self.assertEqual(response.headers["X-Frame-Options"], "DENY")


if __name__ == "__main__":
    unittest.main()
